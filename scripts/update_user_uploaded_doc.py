# update_user_uploaded_doc.py
import docx, os, sys
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_in = r"C:\CANHBAOTHIENTAI\docs\user_uploaded_tieu_luan.docx"
doc = docx.Document(doc_in)

def insert_figure_after(target_p, image_path, caption_text, width=Inches(5.8), intro_text=None):
    elements = []
    if intro_text:
        p_intro = doc.add_paragraph()
        p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_intro.paragraph_format.line_spacing = 1.5
        p_intro.paragraph_format.space_after = Pt(4)
        r = p_intro.add_run(intro_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        elements.append(p_intro._element)

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run = p_img.add_run()
    run.add_picture(image_path, width=width)
    elements.append(p_img._element)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(100, 100, 100)
    elements.append(p_cap._element)

    cur = target_p._element
    for elem in elements:
        cur.addnext(elem)
        cur = elem

img_dir = r"C:\CANHBAOTHIENTAI\docs\images"
img_nguyen_ly = os.path.join(img_dir, "so_do_nguyen_ly_hoat_dong.jpg")
img_lcd = os.path.join(img_dir, "lcd1602_module.jpg")
img_khoi = os.path.join(img_dir, "so_do_khoi_he_thong.jpg")
img_breadboard = os.path.join(img_dir, "so_do_dau_noi_breadboard.jpg")

# 1. Mục 5: Sau Bước 5
for p in doc.paragraphs:
    if "Bước 5 - Máy chủ VPS tiếp nhận" in p.text:
        insert_figure_after(p, img_nguyen_ly, "Hình 0.2. Sơ đồ mô tả dòng chảy dữ liệu và nguyên lý hoạt động tổng thể của hệ thống", width=Inches(5.8))
        print("Inserted Hinh 0.2 after Buoc 5")
        break

# 2. Mục 2.6: Sau đoạn mô tả LCD 1602
for p in doc.paragraphs:
    if "Màn hình tinh thể lỏng LCD 1602 có khả năng hiển thị 2 dòng" in p.text:
        insert_figure_after(p, img_lcd, "Hình 2.6. Màn hình tinh thể lỏng LCD 1602 kết hợp module chuyển đổi I2C PCF8574", width=Inches(3.8))
        print("Inserted Hinh 2.6 after LCD 1602")
        break

# 3. Mục 3.1: Sau đoạn Khối Truyền thông Đám mây
for p in doc.paragraphs:
    if "Khối Truyền thông Đám mây" in p.text and "FastAPI" in p.text:
        insert_figure_after(p, img_khoi, "Hình 3.1. Sơ đồ khối kiến trúc phần cứng tổng thể của Hệ thống Cảnh báo Thiên tai đa năng", width=Inches(5.6))
        print("Inserted Hinh 3.1 after So do khoi")
        break

# 4. Mục 3.3: Sau bảng nối chân, cập nhật Hình 3.1 thành Hình 3.2, và chèn Hình 3.3
for p in doc.paragraphs:
    if "Hình 3.1. Sơ đồ nối chân phần cứng chi tiết của Hệ thống Cảnh báo" in p.text:
        p.text = "Hình 3.2. Sơ đồ nối chân phần cứng chi tiết của Hệ thống Cảnh báo Thiên tai đa năng ESP32"
        insert_figure_after(p, img_breadboard, "Hình 3.3. Sơ đồ nguyên lý đấu nối dây các linh kiện phần cứng trên Breadboard", width=Inches(5.8), intro_text="Bên cạnh bảng nối chân, mô hình đấu nối dây thực tế trên Breadboard được mô phỏng trực quan giúp quá trình lắp ráp mạch diễn ra chuẩn xác, tránh ngắn mạch:")
        print("Updated Hinh 3.2 and inserted Hinh 3.3 after Pinout")
        break

# 5. Cập nhật Hình 3.2 cũ thành Hình 3.4
for p in doc.paragraphs:
    if "Hình 3.2. Biểu đồ so sánh độ trễ" in p.text:
        p.text = "Hình 3.4. Biểu đồ so sánh độ trễ kích hoạt còi tại chỗ trước và sau khi ứng dụng FreeRTOS đa nhân"
        print("Updated Hinh 3.4 for latency chart")
        break

# 6. Phụ lục B: Chèn Hình B.2 sau Hình B.1
for p in doc.paragraphs:
    if "Hình B.1. Sơ đồ nối dây phần cứng chi tiết" in p.text:
        insert_figure_after(p, img_breadboard, "Hình B.2. Sơ đồ bố trí và đấu nối dây mô hình phần cứng thực nghiệm trên Breadboard", width=Inches(5.8))
        print("Inserted Hinh B.2 after Phu luc B.1")
        break

# Lưu đè lên các vị trí chuẩn
out_docs = r"C:\CANHBAOTHIENTAI\docs\user_uploaded_tieu_luan_updated.docx"
out_desktop_1 = r"C:\Users\Admin\Desktop\BaoCao_TieuLuan_GDU_MauChuan_PhanCongDong.docx"
out_desktop_2 = r"C:\Users\Admin\Desktop\BaoCao_CuoiKy_DeTai5_CanhBaoThienTai_PhanCongDong.docx"
out_docs_1 = r"C:\CANHBAOTHIENTAI\docs\BaoCao_TieuLuan_GDU_MauChuan_PhanCongDong.docx"
out_docs_2 = r"C:\CANHBAOTHIENTAI\docs\BaoCao_CuoiKy_DeTai5_CanhBaoThienTai_PhanCongDong.docx"
out_downloads = r"C:\Users\Admin\Downloads\FinalExam_DeTai5_CanhBaoThienTai_PhanCongDong_25150018\BaoCao_TieuLuan_GDU_MauChuan_PhanCongDong.docx"

doc.save(out_docs)
for p in [out_desktop_1, out_desktop_2, out_docs_1, out_docs_2, out_downloads]:
    try:
        doc.save(p)
        print("Saved successfully to:", p)
    except Exception as e:
        print("Error saving to:", p, e)

# Chuyển đổi và lưu cả file .doc gốc nếu người dùng cần định dạng .doc
import win32com.client
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
doc_temp = word.Documents.Open(os.path.abspath(out_docs))
out_doc_desktop = r"C:\Users\Admin\Desktop\BaoCao_TieuLuan_GDU_MauChuan.doc"
doc_temp.SaveAs2(os.path.abspath(out_doc_desktop), FileFormat=0) # 0 = wdFormatDocument (.doc)
doc_temp.Close()
word.Quit()
print("Saved .doc format to:", out_doc_desktop)
print("ALL DONE!")
